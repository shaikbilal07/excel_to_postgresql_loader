import argparse
import datetime as dt
import json
import logging
import logging.handlers
import multiprocessing
import os
import platform
import re
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from contextlib import contextmanager
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation
from functools import partial
from itertools import chain, islice
from pathlib import Path

import psycopg
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from psycopg import sql
from python_calamine import CalamineWorkbook, SheetTypeEnum


ReadOnlyWorksheet._get_size = lambda self: None

log = logging.getLogger("task_one")

TWO_TB = 2 * 10**12
CHECKPOINT_TABLE = "excel_load_checkpoint"
SAMPLE_ROWS = 1000
REJECTS_LOGGED_PER_SHEET = 5
INTEGER_LIMITS = {"int2": 2**15, "int4": 2**31, "int8": 2**63}
TEXT_TYPES = {"text", "varchar", "bpchar", "name", "citext"}
TRUE_VALUES = {"true", "t", "yes", "y", "1"}
FALSE_VALUES = {"false", "f", "no", "n", "0"}

COLUMNS_QUERY = """
    SELECT a.attname,
           t.typname,
           a.attnotnull,
           a.atthasdef OR a.attidentity <> '' OR a.attgenerated <> '',
           CASE WHEN t.typname IN ('varchar', 'bpchar') AND a.atttypmod > 4
                THEN a.atttypmod - 4 END
    FROM pg_attribute a
    JOIN pg_type t ON t.oid = a.atttypid
    WHERE a.attrelid = to_regclass(%s) AND a.attnum > 0 AND NOT a.attisdropped
    ORDER BY a.attnum
"""

CREATE_CHECKPOINT_TABLE = """
    CREATE TABLE IF NOT EXISTS {} (
        target_table text NOT NULL,
        file_path text NOT NULL,
        sheet_name text NOT NULL,
        rows_loaded bigint NOT NULL,
        rows_rejected bigint NOT NULL,
        loaded_at timestamptz NOT NULL DEFAULT now(),
        PRIMARY KEY (target_table, file_path, sheet_name)
    )
"""

SAVE_CHECKPOINT = """
    INSERT INTO {} (target_table, file_path, sheet_name, rows_loaded, rows_rejected)
    VALUES (%s, %s, %s, %s, %s)
    ON CONFLICT (target_table, file_path, sheet_name) DO UPDATE
    SET rows_loaded = EXCLUDED.rows_loaded,
        rows_rejected = EXCLUDED.rows_rejected,
        loaded_at = now()
"""


class SchemaError(ValueError):
    pass


@dataclass(frozen=True)
class Column:
    name: str
    type_name: str
    not_null: bool = False
    has_default: bool = False
    max_length: int | None = None


@dataclass
class LoadSettings:
    database: dict
    table: tuple[str, ...]
    columns: tuple[Column, ...]
    batch_size: int
    engine: str = "openpyxl"
    retries: int = 3
    retry_delay: float = 2.0


@dataclass(frozen=True)
class WorkbookTask:
    path: Path
    key: str
    size: int
    completed_sheets: frozenset[str] = frozenset()


@dataclass
class SheetResult:
    sheet: str
    status: str = "loaded"
    rows_loaded: int = 0
    rows_rejected: int = 0
    batches: int = 0
    failed_batches: int = 0
    seconds: float = 0.0
    error: str = ""


@dataclass
class WorkbookResult:
    key: str
    size: int
    sheets: list[SheetResult] = field(default_factory=list)
    error: str = ""
    seconds: float = 0.0


@dataclass
class RunTotals:
    files_total: int = 0
    files_processed: int = 0
    files_skipped: int = 0
    files_failed: int = 0
    sheets_processed: int = 0
    sheets_skipped: int = 0
    sheets_failed: int = 0
    rows_loaded: int = 0
    rows_rejected: int = 0
    batches: int = 0
    failed_batches: int = 0
    bytes_processed: int = 0
    bytes_seen: int = 0

    def add(self, result: WorkbookResult) -> None:
        self.bytes_seen += result.size
        for sheet in result.sheets:
            if sheet.status == "skipped":
                self.sheets_skipped += 1
            elif sheet.status == "failed":
                self.sheets_failed += 1
            else:
                self.sheets_processed += 1
            self.rows_loaded += sheet.rows_loaded
            self.rows_rejected += sheet.rows_rejected
            self.batches += sheet.batches
            self.failed_batches += sheet.failed_batches
        if result.error:
            self.files_failed += 1
        elif result.sheets and all(s.status == "skipped" for s in result.sheets):
            self.files_skipped += 1
        else:
            self.files_processed += 1
            self.bytes_processed += result.size


@dataclass
class BenchmarkSummary:
    started_at: dt.datetime
    finished_at: dt.datetime
    seconds: float
    totals: RunTotals
    table: str
    engine: str
    workers: int
    batch_size: int
    table_growth: int | None = None
    interrupted: bool = False

    @property
    def bytes_per_second(self) -> float:
        return rate(self.totals.bytes_processed, self.seconds)

    @property
    def rows_per_second(self) -> float:
        return rate(self.totals.rows_loaded, self.seconds)


def positive_int(text: str) -> int:
    value = int(text)
    if value < 1:
        raise argparse.ArgumentTypeError("must be 1 or more")
    return value


def non_negative_int(text: str) -> int:
    value = int(text)
    if value < 0:
        raise argparse.ArgumentTypeError("must be 0 or more")
    return value


def table_name(text: str) -> tuple[str, ...]:
    parts = tuple(part.strip() for part in text.split("."))
    if len(parts) > 2 or not all(parts):
        raise argparse.ArgumentTypeError("expected TABLE or SCHEMA.TABLE")
    return parts


def default_workers() -> int:
    return max(1, (os.cpu_count() or 2) // 2)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Load every worksheet of every Excel workbook in a directory into one "
            "PostgreSQL table with COPY and report the measured execution time."
        ),
        epilog=(
            "The password is never accepted on the command line: set PGPASSWORD or "
            "use a pgpass file."
        ),
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        required=True,
        help="directory with the Excel workbooks (subdirectories are included)",
    )
    parser.add_argument(
        "--table",
        type=table_name,
        required=True,
        help="target table as TABLE or SCHEMA.TABLE; created from the data if missing",
    )
    parser.add_argument(
        "--pattern",
        default="*.xlsx",
        help="workbook file name pattern (default: %(default)s)",
    )
    parser.add_argument(
        "--workers",
        type=positive_int,
        default=default_workers(),
        help="worker processes, each loading one workbook at a time "
        "(default: half the logical CPUs, here %(default)s)",
    )
    parser.add_argument(
        "--batch-size",
        type=positive_int,
        default=50_000,
        help="rows per COPY batch (default: %(default)s)",
    )
    parser.add_argument(
        "--engine",
        choices=("openpyxl", "calamine"),
        default="openpyxl",
        help="Excel reader: openpyxl streams rows with flat memory use; calamine is "
        "much faster but holds one worksheet in memory at a time "
        "(default: %(default)s)",
    )
    parser.add_argument(
        "--retries",
        type=non_negative_int,
        default=3,
        help="retries per sheet after a lost connection or server error "
        "(default: %(default)s)",
    )
    parser.add_argument(
        "--truncate",
        action="store_true",
        help="empty the target table and its load checkpoints before loading",
    )
    parser.add_argument(
        "--unlogged",
        action="store_true",
        help="create a missing target table as UNLOGGED (faster, not crash-safe)",
    )
    parser.add_argument(
        "--report-file", type=Path, help="also write the benchmark summary as JSON"
    )
    parser.add_argument("--log-file", type=Path, help="also write the log to a file")
    parser.add_argument("--verbose", action="store_true", help="log every batch")

    database = parser.add_argument_group("PostgreSQL connection")
    database.add_argument(
        "--host",
        default=os.environ.get("PGHOST", "localhost"),
        help="server host (default: $PGHOST or localhost)",
    )
    database.add_argument(
        "--port",
        type=int,
        default=os.environ.get("PGPORT", "5432"),
        help="server port (default: $PGPORT or 5432)",
    )
    database.add_argument(
        "--dbname",
        default=os.environ.get("PGDATABASE", "postgres"),
        help="database name (default: $PGDATABASE or postgres)",
    )
    database.add_argument(
        "--user",
        default=os.environ.get("PGUSER", "postgres"),
        help="user name (default: $PGUSER or postgres)",
    )

    args = parser.parse_args(argv)
    if sys.platform == "win32" and args.workers > 61:
        parser.error("--workers cannot be more than 61 on Windows")
    return args


def discover_excel_files(input_dir: Path, pattern: str = "*.xlsx") -> list[Path]:
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input directory not found: {input_dir}")
    return sorted(
        path
        for path in input_dir.rglob(pattern)
        if path.is_file() and not path.name.startswith("~$")
    )


@contextmanager
def open_workbook(path: Path, engine: str):
    if engine == "calamine":
        workbook = CalamineWorkbook.from_path(str(path))
        try:
            yield [
                (sheet.name, partial(calamine_rows, workbook, sheet.name))
                for sheet in workbook.sheets_metadata
                if sheet.typ == SheetTypeEnum.WorkSheet
            ]
        finally:
            workbook.close()
    else:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            yield [
                (sheet.title, partial(openpyxl_rows, sheet))
                for sheet in workbook.worksheets
            ]
        finally:
            workbook.close()


def openpyxl_rows(worksheet):
    # Many writers omit or misreport the sheet dimension; trusting it truncates rows.
    worksheet.reset_dimensions()
    return worksheet.iter_rows(values_only=True)


def calamine_rows(workbook, name: str):
    sheet = workbook.get_sheet_by_name(name)
    # python-calamine panics instead of yielding nothing for a sheet without cells.
    if sheet.start is not None:
        yield from sheet.iter_rows()


def read_sample(files: list[Path], engine: str, limit: int = SAMPLE_ROWS):
    for path in files:
        try:
            with open_workbook(path, engine) as sheets:
                for _, read_rows in sheets:
                    rows = (row for row in read_rows() if not is_blank_row(row))
                    header = next(rows, None)
                    if header is not None:
                        return header, list(islice(rows, limit))
        except Exception as exc:
            log.warning("Cannot sample %s: %s", path, describe(exc))
    return None, []


def is_blank_row(row) -> bool:
    return row.count(None) + row.count("") == len(row)


def normalize_column_name(value: str) -> str:
    name = re.sub(r"\W+", "_", value.strip().lower()).strip("_")
    return name.encode()[:63].decode(errors="ignore")


def read_header(header) -> tuple[int, list[str]]:
    cells = list(header)
    while cells and (cells[-1] is None or cells[-1] == ""):
        cells.pop()
    offset = 0
    while offset < len(cells) and (cells[offset] is None or cells[offset] == ""):
        offset += 1

    names = []
    for position in range(offset, len(cells)):
        cell = cells[position]
        letter = get_column_letter(position + 1)
        if cell is None or cell == "":
            raise SchemaError(f"column {letter} has no header")
        name = normalize_column_name(to_text(cell))
        if not name:
            raise SchemaError(f"column {letter} header {cell!r} is not a usable name")
        if name in names:
            raise SchemaError(f"column {letter} header {cell!r} is a duplicate")
        names.append(name)
    return offset, names


def match_header(header, columns) -> tuple[int, list[Column]]:
    offset, names = read_header(header)
    by_name = {normalize_column_name(column.name): column for column in columns}
    unknown = [name for name in names if name not in by_name]
    if unknown:
        raise SchemaError("columns not in the target table: " + ", ".join(unknown))
    missing = [
        column.name
        for key, column in by_name.items()
        if column.not_null and not column.has_default and key not in names
    ]
    if missing:
        raise SchemaError("missing required columns: " + ", ".join(missing))
    return offset, [by_name[name] for name in names]


def infer_column_type(values) -> str:
    kinds = set()
    for value in values:
        if value is None or value == "":
            continue
        if isinstance(value, bool):
            kinds.add("boolean")
        elif isinstance(value, (int, float)):
            whole = isinstance(value, int) or value.is_integer()
            kinds.add("bigint" if whole else "double precision")
        elif isinstance(value, (dt.datetime, dt.date)):
            kinds.add("timestamp")
        elif isinstance(value, dt.time):
            kinds.add("time")
        elif isinstance(value, dt.timedelta):
            kinds.add("interval")
        else:
            kinds.add("text")
    if kinds == {"bigint", "double precision"}:
        return "double precision"
    return kinds.pop() if len(kinds) == 1 else "text"


def infer_table_definition(header, sample_rows) -> list[tuple[str, str]]:
    offset, names = read_header(header)
    definition = []
    for index, name in enumerate(names, start=offset):
        values = [row[index] for row in sample_rows if index < len(row)]
        definition.append((name, infer_column_type(values)))
    return definition


def to_integer(value, limit: int) -> int:
    if isinstance(value, bool):
        raise ValueError("expected a whole number, got a boolean")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"{value!r} is not a whole number")
        value = int(value)
    elif isinstance(value, str):
        value = int(value)
    elif not isinstance(value, int):
        raise ValueError(f"expected a whole number, got {type(value).__name__}")
    if not -limit <= value < limit:
        raise ValueError(f"{value} is out of range")
    return value


def to_number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise ValueError(f"expected a number, got {type(value).__name__}")
    if isinstance(value, str):
        try:
            return Decimal(value)
        except InvalidOperation:
            raise ValueError(f"{value!r} is not a number") from None
    return value


def to_boolean(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        text = value.strip().lower()
        if text in TRUE_VALUES:
            return True
        if text in FALSE_VALUES:
            return False
    raise ValueError(f"{value!r} is not a boolean")


def to_timestamp(value) -> dt.datetime:
    if isinstance(value, dt.datetime):
        return value
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day)
    if isinstance(value, str):
        return dt.datetime.fromisoformat(value.strip())
    raise ValueError(f"expected a date and time, got {type(value).__name__}")


def to_date(value) -> dt.date:
    if isinstance(value, dt.datetime):
        if value.time() != dt.time():
            raise ValueError(f"{value} has a time of day")
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        return dt.date.fromisoformat(value.strip())
    raise ValueError(f"expected a date, got {type(value).__name__}")


def to_time(value) -> dt.time:
    if isinstance(value, dt.time):
        return value
    if isinstance(value, str):
        return dt.time.fromisoformat(value.strip())
    raise ValueError(f"expected a time of day, got {type(value).__name__}")


def to_text(value, max_length: int | None = None) -> str:
    if isinstance(value, str):
        text = value
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    if "\x00" in text:
        raise ValueError("text contains a NUL character")
    if max_length is not None and len(text) > max_length:
        raise ValueError(f"text is longer than {max_length} characters")
    return text


def pass_through(value):
    return value


CONVERTERS = {
    "float4": to_number,
    "float8": to_number,
    "numeric": to_number,
    "bool": to_boolean,
    "date": to_date,
    "timestamp": to_timestamp,
    "timestamptz": to_timestamp,
    "time": to_time,
}


def column_converter(column: Column):
    if column.type_name in INTEGER_LIMITS:
        convert = partial(to_integer, limit=INTEGER_LIMITS[column.type_name])
    elif column.type_name in TEXT_TYPES:
        convert = partial(to_text, max_length=column.max_length)
    else:
        convert = CONVERTERS.get(column.type_name, pass_through)

    def converter(value):
        if value is None or value == "":
            if column.not_null:
                raise ValueError(f"{column.name}: a value is required")
            return None
        try:
            return convert(value)
        except (ValueError, OverflowError) as exc:
            raise ValueError(f"{column.name}: {exc}") from None

    return converter


def convert_row(row, offset: int, converters) -> list:
    end = offset + len(converters)
    if len(row) > end and not is_blank_row(row[end:]):
        raise ValueError("row has values in columns without a header")
    if offset and not is_blank_row(row[:offset]):
        raise ValueError("row has values in columns without a header")
    cells = row[offset:end]
    if len(cells) < len(converters):
        cells = list(cells) + [None] * (len(converters) - len(cells))
    return [convert(cell) for convert, cell in zip(converters, cells)]


def validated_rows(numbered_rows, offset, converters, label, result: SheetResult):
    for row_number, row in numbered_rows:
        if is_blank_row(row):
            continue
        try:
            yield convert_row(row, offset, converters)
        except ValueError as exc:
            result.rows_rejected += 1
            if result.rows_rejected <= REJECTS_LOGGED_PER_SHEET:
                log.warning("%s: row %d rejected: %s", label, row_number, exc)


def iter_batches(items, size: int):
    # Each batch is a lazy view of the shared iterator, so it has to be consumed
    # completely before the next batch is requested.
    if size < 1:
        raise ValueError("batch size must be 1 or more")
    iterator = iter(items)
    for first in iterator:
        yield chain((first,), islice(iterator, size - 1))


def connection_settings(args: argparse.Namespace) -> dict:
    return {
        "host": args.host,
        "port": args.port,
        "dbname": args.dbname,
        "user": args.user,
        "connect_timeout": 30,
        "application_name": "task_one",
        # A crash can only lose the most recent commits, and those sheets lose
        # their checkpoint with them, so the next run simply loads them again.
        "options": "-c synchronous_commit=off",
    }


def connect(database: dict):
    return psycopg.connect(**database)


def close_connection(connection) -> None:
    if connection is not None:
        connection.close()


def checkpoint_table(table: tuple[str, ...]) -> sql.Identifier:
    return sql.Identifier(*table[:-1], CHECKPOINT_TABLE)


def fetch_columns(connection, table: tuple[str, ...]) -> list[Column]:
    rows = connection.execute(COLUMNS_QUERY, (sql.Identifier(*table).as_string(),))
    return [Column(*row) for row in rows.fetchall()]


def create_table(connection, table, definition, unlogged: bool) -> None:
    connection.execute(
        sql.SQL("CREATE {}TABLE {} ({})").format(
            sql.SQL("UNLOGGED " if unlogged else ""),
            sql.Identifier(*table),
            sql.SQL(", ").join(
                sql.SQL("{} {}").format(sql.Identifier(name), sql.SQL(kind))
                for name, kind in definition
            ),
        )
    )


def prepare_target(connection, args: argparse.Namespace, files) -> list[Column]:
    label = ".".join(args.table)
    connection.execute(
        sql.SQL(CREATE_CHECKPOINT_TABLE).format(checkpoint_table(args.table))
    )
    columns = fetch_columns(connection, args.table)
    if not columns:
        header, sample = read_sample(files, args.engine)
        if header is None:
            raise SchemaError("no worksheet with a header row was found in the input")
        definition = infer_table_definition(header, sample)
        create_table(connection, args.table, definition, args.unlogged)
        columns = fetch_columns(connection, args.table)
        log.info(
            "Created %stable %s (%s)",
            "UNLOGGED " if args.unlogged else "",
            label,
            ", ".join(f"{name} {kind}" for name, kind in definition),
        )
    elif args.unlogged:
        log.info("Table %s already exists, so --unlogged has no effect", label)

    if args.truncate:
        connection.execute(sql.SQL("TRUNCATE {}").format(sql.Identifier(*args.table)))
        connection.execute(
            sql.SQL("DELETE FROM {} WHERE target_table = %s").format(
                checkpoint_table(args.table)
            ),
            (label,),
        )
        log.warning("Truncated %s and cleared its load checkpoints", label)
    connection.commit()
    return columns


def completed_sheets(connection, table: tuple[str, ...]) -> dict[str, set[str]]:
    query = sql.SQL(
        "SELECT file_path, sheet_name FROM {} WHERE target_table = %s"
    ).format(checkpoint_table(table))
    done: dict[str, set[str]] = {}
    for file_path, sheet_name in connection.execute(query, (".".join(table),)):
        done.setdefault(file_path, set()).add(sheet_name)
    return done


def table_size(connection, table: tuple[str, ...]) -> int | None:
    row = connection.execute(
        "SELECT pg_total_relation_size(to_regclass(%s))",
        (sql.Identifier(*table).as_string(),),
    ).fetchone()
    return row[0]


def load_sheet(connection, rows, settings: LoadSettings, file_key, result) -> None:
    label = f"{file_key} [{result.sheet}]"
    started = time.perf_counter()
    numbered_rows = enumerate(rows, start=1)
    header = next((row for _, row in numbered_rows if not is_blank_row(row)), None)
    if header is None:
        result.status = "empty"
        log.info("%s: empty sheet, nothing to load", label)
        return

    offset, columns = match_header(header, settings.columns)
    converters = [column_converter(column) for column in columns]
    copy_statement = sql.SQL("COPY {} ({}) FROM STDIN").format(
        sql.Identifier(*settings.table),
        sql.SQL(", ").join(sql.Identifier(column.name) for column in columns),
    )
    rows_to_load = validated_rows(numbered_rows, offset, converters, label, result)

    # The checkpoint is written in the same transaction as the sheet's rows, so a
    # sheet is either fully loaded and recorded or not loaded at all.
    with connection.cursor() as cursor:
        for batch in iter_batches(rows_to_load, settings.batch_size):
            result.batches += 1
            batch_rows = 0
            with cursor.copy(copy_statement) as copy:
                for values in batch:
                    copy.write_row(values)
                    batch_rows += 1
            result.rows_loaded += batch_rows
            log.debug("%s: batch %d sent, %d rows", label, result.batches, batch_rows)
        cursor.execute(
            sql.SQL(SAVE_CHECKPOINT).format(checkpoint_table(settings.table)),
            (
                ".".join(settings.table),
                file_key,
                result.sheet,
                result.rows_loaded,
                result.rows_rejected,
            ),
        )
    connection.commit()

    result.seconds = time.perf_counter() - started
    if result.rows_rejected > REJECTS_LOGGED_PER_SHEET:
        log.warning("%s: %d rows rejected in total", label, result.rows_rejected)
    log.info(
        "%s: loaded %s rows in %d batches, %.1fs (%s rows/s)",
        label,
        f"{result.rows_loaded:,}",
        result.batches,
        result.seconds,
        f"{rate(result.rows_loaded, result.seconds):,.0f}",
    )


def describe(error: BaseException) -> str:
    message = " ".join(str(error).split())
    return f"{type(error).__name__}: {message}" if message else type(error).__name__


def mark_failed(result: SheetResult, error: Exception, label: str) -> None:
    result.status = "failed"
    result.error = describe(error)
    result.failed_batches = result.batches
    result.batches = result.rows_loaded = result.rows_rejected = 0
    unexpected = not isinstance(error, (psycopg.Error, SchemaError))
    log.error("%s: sheet not loaded: %s", label, result.error, exc_info=unexpected)


def process_workbook(task: WorkbookTask, settings: LoadSettings) -> WorkbookResult:
    started = time.perf_counter()
    outcome = WorkbookResult(task.key, task.size)
    connection = None
    try:
        with open_workbook(task.path, settings.engine) as sheets:
            for sheet_name, read_rows in sheets:
                label = f"{task.key} [{sheet_name}]"
                if sheet_name in task.completed_sheets:
                    log.debug("%s: loaded by an earlier run, skipping", label)
                    outcome.sheets.append(SheetResult(sheet_name, status="skipped"))
                    continue

                for attempt in range(settings.retries + 1):
                    result = SheetResult(sheet_name)
                    try:
                        if connection is None:
                            connection = connect(settings.database)
                        load_sheet(connection, read_rows(), settings, task.key, result)
                        break
                    except Exception as exc:
                        close_connection(connection)
                        connection = None
                        retryable = isinstance(exc, psycopg.OperationalError)
                        if not retryable or attempt == settings.retries:
                            mark_failed(result, exc, label)
                            break
                        delay = settings.retry_delay * 2**attempt
                        log.warning(
                            "%s: %s; retrying in %.0fs (attempt %d of %d)",
                            label,
                            describe(exc),
                            delay,
                            attempt + 2,
                            settings.retries + 1,
                        )
                        time.sleep(delay)
                outcome.sheets.append(result)
    except Exception as exc:
        outcome.error = describe(exc)
        log.error("%s: cannot read workbook: %s", task.key, outcome.error)
    finally:
        close_connection(connection)
    outcome.seconds = time.perf_counter() - started
    return outcome


def load_workbooks(tasks, settings: LoadSettings, workers, log_queue, log_level):
    if workers == 1:
        for task in tasks:
            yield process_workbook(task, settings)
        return

    pool = ProcessPoolExecutor(
        max_workers=workers,
        mp_context=multiprocessing.get_context("spawn"),
        initializer=use_queue_logging,
        initargs=(log_queue, log_level),
    )
    try:
        futures = {
            pool.submit(process_workbook, task, settings): task for task in tasks
        }
        for future in as_completed(futures):
            task = futures[future]
            try:
                yield future.result()
            except Exception as exc:
                log.error("%s: worker failed: %s", task.key, describe(exc))
                yield WorkbookResult(task.key, task.size, error=describe(exc))
    finally:
        pool.shutdown(wait=True, cancel_futures=True)


def rate(amount: float, seconds: float) -> float:
    return amount / seconds if seconds > 0 else 0.0


def format_bytes(count: float) -> str:
    value = float(count)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if abs(value) < 1000:
            return f"{value:,.2f} {unit}"
        value /= 1000
    return f"{value:,.2f} PB"


def format_duration(seconds: float) -> str:
    minutes, secs = divmod(max(seconds, 0.0), 60)
    hours, minutes = divmod(int(minutes), 60)
    days, hours = divmod(hours, 24)
    text = f"{hours:02d}h {minutes:02d}m {secs:05.2f}s"
    return f"{days}d {text}" if days else text


def log_progress(result: WorkbookResult, totals: RunTotals, total_bytes, elapsed):
    finished = totals.files_processed + totals.files_skipped + totals.files_failed
    speed = rate(totals.bytes_processed, elapsed)
    remaining = total_bytes - totals.bytes_seen
    eta = format_duration(remaining / speed) if speed else "n/a"
    log.info(
        "[%d/%d] %s: %s rows in %.1fs | total %s rows, %s, %s/s, %s rows/s, ETA %s",
        finished,
        totals.files_total,
        result.key,
        f"{sum(sheet.rows_loaded for sheet in result.sheets):,}",
        result.seconds,
        f"{totals.rows_loaded:,}",
        format_bytes(totals.bytes_processed),
        format_bytes(speed),
        f"{rate(totals.rows_loaded, elapsed):,.0f}",
        eta,
    )


def projection_lines(summary: BenchmarkSummary) -> list[str]:
    totals = summary.totals
    if totals.bytes_processed >= TWO_TB:
        return [
            "RESULT: ACTUAL 2 TB RUN",
            f"Measured on {format_bytes(totals.bytes_processed)} of Excel data; the "
            "execution time above is the actual 2 TB result.",
        ]
    lines = [
        "RESULT: ACTUAL for the dataset above. 2 TB figures below are PROJECTED.",
        f"Tested Dataset    : {format_bytes(totals.bytes_processed)} "
        f"({totals.bytes_processed / TWO_TB:.4%} of 2 TB) - Actual",
    ]
    speed = summary.bytes_per_second
    if speed <= 0:
        return lines + ["Projected 2 TB    : not available, no data was processed"]
    lines.append(
        f"2 TB Load Time    : {format_duration(TWO_TB / speed)} "
        f"at {format_bytes(speed)}/s - Projected"
    )
    scale = TWO_TB / totals.bytes_processed
    lines.append(f"2 TB Row Count    : ~{totals.rows_loaded * scale:,.0f} - Projected")
    if summary.table_growth is not None and summary.table_growth > 0:
        lines.append(
            f"2 TB Table Size   : ~{format_bytes(summary.table_growth * scale)} "
            "in PostgreSQL - Projected"
        )
    lines.append("Projections assume the measured throughput stays constant.")
    return lines


def format_summary(summary: BenchmarkSummary) -> str:
    totals = summary.totals
    growth = summary.table_growth
    bar = "=" * 64
    title = "EXCEL TO POSTGRESQL BENCHMARK"
    if summary.interrupted:
        title += " (INTERRUPTED, PARTIAL RUN)"
    lines = [
        bar,
        title,
        bar,
        f"Start Time        : {summary.started_at.isoformat(' ', 'seconds')}",
        f"End Time          : {summary.finished_at.isoformat(' ', 'seconds')}",
        f"Execution Time    : {format_duration(summary.seconds)} "
        f"({summary.seconds:,.2f} seconds)",
        f"Files Processed   : {totals.files_processed:,} of {totals.files_total:,}",
        f"Sheets Processed  : {totals.sheets_processed:,}",
        f"Rows Processed    : {totals.rows_loaded:,}",
        f"Data Processed    : {format_bytes(totals.bytes_processed)} "
        f"({totals.bytes_processed:,} bytes)",
        f"Throughput        : {format_bytes(summary.bytes_per_second)}/s",
        f"Rows/Second       : {summary.rows_per_second:,.0f}",
        f"Successful Batches: {totals.batches:,}",
        f"Failed Batches    : {totals.failed_batches:,}",
        "-" * 64,
        f"Rows Rejected     : {totals.rows_rejected:,}",
        f"Sheets Failed     : {totals.sheets_failed:,}",
        f"Sheets Skipped    : {totals.sheets_skipped:,} (loaded by an earlier run)",
        f"Files Failed      : {totals.files_failed:,}",
        f"Target Table      : {summary.table}",
        f"Table Growth      : {format_bytes(growth) if growth is not None else 'n/a'}",
        f"Engine            : {summary.engine}",
        f"Workers           : {summary.workers}",
        f"Batch Size        : {summary.batch_size:,} rows",
        bar,
        *projection_lines(summary),
        bar,
    ]
    return "\n".join(lines)


def summary_record(summary: BenchmarkSummary) -> dict:
    totals = summary.totals
    projected = (
        TWO_TB / summary.bytes_per_second if summary.bytes_per_second > 0 else None
    )
    return {
        "result": "actual",
        "interrupted": summary.interrupted,
        "start_time": summary.started_at.isoformat(),
        "end_time": summary.finished_at.isoformat(),
        "execution_seconds": round(summary.seconds, 6),
        **asdict(totals),
        "throughput_bytes_per_second": round(summary.bytes_per_second, 1),
        "rows_per_second": round(summary.rows_per_second, 1),
        "table": summary.table,
        "table_growth_bytes": summary.table_growth,
        "engine": summary.engine,
        "workers": summary.workers,
        "batch_size": summary.batch_size,
        "full_2tb_processed": totals.bytes_processed >= TWO_TB,
        "projected_2tb_seconds": None if projected is None else round(projected, 1),
        "machine": {
            "platform": platform.platform(),
            "python": platform.python_version(),
            "logical_cpus": os.cpu_count(),
        },
    }


def run(args: argparse.Namespace, log_queue) -> int:
    started_at = dt.datetime.now().astimezone()
    clock = time.perf_counter()
    table_label = ".".join(args.table)

    files = discover_excel_files(args.input_dir, args.pattern)
    if not files:
        log.error("No files matching %r found in %s", args.pattern, args.input_dir)
        return 1

    database = connection_settings(args)
    with connect(database) as connection:
        columns = prepare_target(connection, args, files)
        done = completed_sheets(connection, args.table)
        size_before = table_size(connection, args.table)

    tasks = []
    for path in files:
        key = path.relative_to(args.input_dir).as_posix()
        completed = frozenset(done.get(key, ()))
        tasks.append(WorkbookTask(path, key, path.stat().st_size, completed))
    total_bytes = sum(task.size for task in tasks)
    settings = LoadSettings(
        database=database,
        table=args.table,
        columns=tuple(columns),
        batch_size=args.batch_size,
        engine=args.engine,
        retries=args.retries,
    )
    log.info(
        "Loading %d workbooks (%s) into %s: engine=%s, workers=%d, batch size=%s",
        len(tasks),
        format_bytes(total_bytes),
        table_label,
        args.engine,
        args.workers,
        f"{args.batch_size:,}",
    )

    totals = RunTotals(files_total=len(tasks))
    interrupted = False
    log_level = logging.getLogger().level
    try:
        for result in load_workbooks(
            tasks, settings, args.workers, log_queue, log_level
        ):
            totals.add(result)
            log_progress(result, totals, total_bytes, time.perf_counter() - clock)
    except KeyboardInterrupt:
        interrupted = True
        log.warning("Interrupted. Committed sheets will be skipped on the next run.")
    seconds = time.perf_counter() - clock
    finished_at = dt.datetime.now().astimezone()

    growth = None
    try:
        with connect(database) as connection:
            size_after = table_size(connection, args.table)
        if size_before is not None and size_after is not None:
            growth = size_after - size_before
    except psycopg.Error as exc:
        log.warning("Could not measure the table size: %s", describe(exc))

    summary = BenchmarkSummary(
        started_at=started_at,
        finished_at=finished_at,
        seconds=seconds,
        totals=totals,
        table=table_label,
        engine=args.engine,
        workers=args.workers,
        batch_size=args.batch_size,
        table_growth=growth,
        interrupted=interrupted,
    )
    log.info("Benchmark finished\n%s", format_summary(summary))
    if args.report_file:
        args.report_file.write_text(
            json.dumps(summary_record(summary), indent=2), encoding="utf-8"
        )
        log.info("Benchmark report written to %s", args.report_file)

    if interrupted:
        return 130
    return 1 if totals.files_failed or totals.sheets_failed else 0


def use_queue_logging(log_queue, level: int) -> None:
    root = logging.getLogger()
    root.handlers[:] = [logging.handlers.QueueHandler(log_queue)]
    root.setLevel(level)
    logging.captureWarnings(True)


def configure_logging(level: int, log_file: Path | None = None):
    formatter = logging.Formatter(
        "%(asctime)s %(levelname)-7s %(processName)s: %(message)s",
        "%Y-%m-%d %H:%M:%S",
    )
    handlers: list[logging.Handler] = [logging.StreamHandler(sys.stdout)]
    if log_file is not None:
        handlers.append(logging.FileHandler(log_file, encoding="utf-8"))
    for handler in handlers:
        handler.setFormatter(formatter)
    log_queue = multiprocessing.get_context("spawn").Queue()
    listener = logging.handlers.QueueListener(log_queue, *handlers)
    use_queue_logging(log_queue, level)
    listener.start()
    return log_queue, listener


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    level = logging.DEBUG if args.verbose else logging.INFO
    log_queue, listener = configure_logging(level, args.log_file)
    try:
        return run(args, log_queue)
    except KeyboardInterrupt:
        log.warning("Interrupted")
        return 130
    except (OSError, psycopg.Error, SchemaError) as exc:
        log.error("%s", describe(exc))
        return 1
    finally:
        listener.stop()


if __name__ == "__main__":
    sys.exit(main())
