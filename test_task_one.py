import argparse
import datetime as dt
import json
import logging
import os
import uuid
from decimal import Decimal

import psycopg
import pytest
from openpyxl import Workbook, load_workbook
from psycopg import sql

import generate_test_data
import task_one
from task_one import (
    BenchmarkSummary,
    Column,
    LoadSettings,
    RunTotals,
    SchemaError,
    SheetResult,
    WorkbookResult,
    WorkbookTask,
)

ENGINES = ["openpyxl", "calamine"]
HEADER = ["order_id", "customer", "amount", "ordered_at"]
ORDER_COLUMNS = (
    Column("order_id", "int8", not_null=True),
    Column("customer", "text"),
    Column("amount", "numeric"),
    Column("ordered_at", "timestamp"),
)

requires_postgres = pytest.mark.skipif(
    os.environ.get("RUN_POSTGRES_TESTS") != "1",
    reason="set RUN_POSTGRES_TESTS=1 and the PG* variables to test against PostgreSQL",
)


def order_row(number):
    return [
        number,
        f"Customer {number}",
        number * 1.5,
        dt.datetime(2024, 1, 1) + dt.timedelta(days=number),
    ]


def write_workbook(path, sheets):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


def make_task(path, completed=()):
    return WorkbookTask(path, path.name, path.stat().st_size, frozenset(completed))


def make_settings(**overrides):
    values = {
        "database": {},
        "table": ("orders",),
        "columns": ORDER_COLUMNS,
        "batch_size": 2,
        "retries": 2,
        "retry_delay": 0,
    }
    values.update(overrides)
    return LoadSettings(**values)


class FakeCopy:
    def __init__(self, connection):
        self.connection = connection
        self.rows = []

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        if exc_type is None:
            self.connection.finish_copy(self.rows)
        return False

    def write_row(self, row):
        self.rows.append(list(row))


class FakeCursor:
    def __init__(self, connection):
        self.connection = connection

    def __enter__(self):
        return self

    def __exit__(self, *exc_info):
        return False

    def copy(self, statement):
        return FakeCopy(self.connection)

    def execute(self, statement, params=None):
        self.connection.pending_checkpoints.append(params)


class FakeConnection:
    def __init__(self, fail_on_copy=None, error=None):
        self.fail_on_copy = fail_on_copy
        self.error = error
        self.copies = 0
        self.pending_batches = []
        self.pending_checkpoints = []
        self.batches = []
        self.checkpoints = []
        self.commits = 0
        self.closed = False

    def __enter__(self):
        return self

    def __exit__(self, *exc_info):
        self.close()
        return False

    def cursor(self):
        return FakeCursor(self)

    def finish_copy(self, rows):
        self.copies += 1
        if self.copies == self.fail_on_copy:
            raise self.error
        self.pending_batches.append(rows)

    def commit(self):
        self.batches += self.pending_batches
        self.checkpoints += self.pending_checkpoints
        self.pending_batches, self.pending_checkpoints = [], []
        self.commits += 1

    def close(self):
        self.pending_batches, self.pending_checkpoints = [], []
        self.closed = True


class FakeDatabase:
    def __init__(self):
        self.connections = []
        self.failures = []

    def connect(self, settings):
        fail_on_copy, error = self.failures.pop(0) if self.failures else (None, None)
        connection = FakeConnection(fail_on_copy, error)
        self.connections.append(connection)
        return connection

    @property
    def rows(self):
        return [
            row
            for connection in self.connections
            for batch in connection.batches
            for row in batch
        ]

    @property
    def checkpoints(self):
        return [
            params
            for connection in self.connections
            for params in connection.checkpoints
        ]


@pytest.fixture
def database(monkeypatch):
    fake = FakeDatabase()
    monkeypatch.setattr(task_one, "connect", fake.connect)
    return fake


@pytest.fixture
def run_main():
    def run(argv):
        root = logging.getLogger()
        handlers, level = root.handlers[:], root.level
        try:
            return task_one.main(argv)
        finally:
            root.handlers[:] = handlers
            root.setLevel(level)
            logging.captureWarnings(False)

    return run


def test_discover_finds_workbooks_recursively_and_skips_lock_files(tmp_path):
    first = write_workbook(tmp_path / "b.xlsx", {"Sheet": [["id"]]})
    second = write_workbook(tmp_path / "nested" / "a.xlsx", {"Sheet": [["id"]]})
    (tmp_path / "~$b.xlsx").write_bytes(b"lock file")
    (tmp_path / "notes.csv").write_text("id\n1\n")

    assert task_one.discover_excel_files(tmp_path) == sorted([first, second])


def test_discover_returns_empty_list_when_there_are_no_workbooks(tmp_path):
    (tmp_path / "data.csv").write_text("id\n1\n")

    assert task_one.discover_excel_files(tmp_path) == []


def test_discover_rejects_missing_directory(tmp_path):
    with pytest.raises(FileNotFoundError):
        task_one.discover_excel_files(tmp_path / "missing")


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("Order ID", "order_id"),
        ("  amount ($) ", "amount"),
        ("Customer-Name", "customer_name"),
        ("ÜNIT Price", "ünit_price"),
    ],
)
def test_normalize_column_name(raw, expected):
    assert task_one.normalize_column_name(raw) == expected


def test_match_header_maps_names_case_insensitively_in_sheet_order():
    offset, columns = task_one.match_header(
        ["Customer", "ORDER_ID", "Ordered At", "Amount"], ORDER_COLUMNS
    )

    assert offset == 0
    assert [column.name for column in columns] == [
        "customer",
        "order_id",
        "ordered_at",
        "amount",
    ]


def test_match_header_ignores_blank_cells_around_the_header():
    offset, columns = task_one.match_header(
        [None, "", "order_id", "customer", None], ORDER_COLUMNS
    )

    assert offset == 2
    assert [column.name for column in columns] == ["order_id", "customer"]


@pytest.mark.parametrize(
    "header, message",
    [
        (["order_id", "discount"], "not in the target table: discount"),
        (["order_id", "Order ID"], "duplicate"),
        (["order_id", None, "amount"], "column B has no header"),
        (["customer", "amount"], "missing required columns: order_id"),
    ],
)
def test_match_header_rejects_invalid_headers(header, message):
    with pytest.raises(SchemaError, match=message):
        task_one.match_header(header, ORDER_COLUMNS)


def test_match_header_allows_missing_columns_with_defaults():
    columns = (
        Column("id", "int8", not_null=True, has_default=True),
        Column("name", "text"),
    )

    _, matched = task_one.match_header(["name"], columns)

    assert [column.name for column in matched] == ["name"]


@pytest.mark.parametrize(
    "column, raw, expected",
    [
        (Column("n", "int8"), 42, 42),
        (Column("n", "int8"), 42.0, 42),
        (Column("n", "int4"), " 17 ", 17),
        (Column("n", "numeric"), 1.25, 1.25),
        (Column("n", "numeric"), "19.99", Decimal("19.99")),
        (Column("n", "float8"), 3, 3),
        (Column("n", "bool"), "Yes", True),
        (Column("n", "bool"), 0, False),
        (Column("n", "timestamp"), dt.date(2024, 5, 1), dt.datetime(2024, 5, 1)),
        (
            Column("n", "timestamp"),
            "2024-05-01 10:30:00",
            dt.datetime(2024, 5, 1, 10, 30),
        ),
        (Column("n", "date"), dt.datetime(2024, 5, 1), dt.date(2024, 5, 1)),
        (Column("n", "date"), "2024-05-01", dt.date(2024, 5, 1)),
        (Column("n", "time"), "08:15:00", dt.time(8, 15)),
        (Column("n", "text"), 12345.0, "12345"),
        (Column("n", "text"), 2.5, "2.5"),
        (Column("n", "varchar", max_length=3), "abc", "abc"),
        (
            Column("n", "uuid"),
            "4b1f8b7e-3c3a-4a5e-9a0b-0f3d2c1b0a99",
            "4b1f8b7e-3c3a-4a5e-9a0b-0f3d2c1b0a99",
        ),
    ],
)
def test_column_converter_accepts_valid_values(column, raw, expected):
    assert task_one.column_converter(column)(raw) == expected


@pytest.mark.parametrize(
    "column, raw",
    [
        (Column("n", "int8"), 1.5),
        (Column("n", "int8"), "12abc"),
        (Column("n", "int8"), True),
        (Column("n", "int2"), 40_000),
        (Column("n", "int8"), float("inf")),
        (Column("n", "numeric"), "twelve"),
        (Column("n", "numeric"), dt.datetime(2024, 1, 1)),
        (Column("n", "bool"), "maybe"),
        (Column("n", "timestamp"), 45000),
        (Column("n", "timestamp"), "next tuesday"),
        (Column("n", "date"), dt.datetime(2024, 5, 1, 12, 30)),
        (Column("n", "varchar", max_length=3), "abcd"),
        (Column("n", "text"), "bad\x00value"),
    ],
)
def test_column_converter_rejects_invalid_values(column, raw):
    with pytest.raises(ValueError, match="^n: "):
        task_one.column_converter(column)(raw)


def test_blank_cells_become_null_unless_the_column_is_required():
    optional = task_one.column_converter(Column("note", "text"))
    required = task_one.column_converter(Column("id", "int8", not_null=True))

    assert optional(None) is None
    assert optional("") is None
    with pytest.raises(ValueError, match="a value is required"):
        required(None)


def test_convert_row_pads_short_rows_and_ignores_trailing_blanks():
    converters = [task_one.column_converter(column) for column in ORDER_COLUMNS]

    assert task_one.convert_row((7, "Ann"), 0, converters) == [7, "Ann", None, None]
    assert task_one.convert_row([7, "Ann", 1.5, None, "", None], 0, converters) == [
        7,
        "Ann",
        1.5,
        None,
    ]


def test_convert_row_rejects_values_outside_the_header():
    converters = [task_one.column_converter(column) for column in ORDER_COLUMNS[:2]]

    with pytest.raises(ValueError, match="without a header"):
        task_one.convert_row((7, "Ann", "extra"), 0, converters)
    with pytest.raises(ValueError, match="without a header"):
        task_one.convert_row(("stray", 7, "Ann"), 1, converters)
    assert task_one.convert_row((None, 7, "Ann"), 1, converters) == [7, "Ann"]


def test_iter_batches_yields_batches_of_the_requested_size():
    batches = [list(batch) for batch in task_one.iter_batches(range(7), 3)]

    assert batches == [[0, 1, 2], [3, 4, 5], [6]]


def test_iter_batches_handles_empty_input_and_rejects_invalid_size():
    assert list(task_one.iter_batches([], 10)) == []
    with pytest.raises(ValueError):
        list(task_one.iter_batches([1, 2], 0))


def test_infer_table_definition_picks_types_from_sample_values():
    header = ["ID", "Price", "Ordered", "Paid", "Note", "Empty"]
    rows = [
        [1, 2.5, dt.datetime(2024, 1, 1), True, "a", None],
        [2.0, 3, dt.date(2024, 1, 2), False, 5, ""],
    ]

    assert task_one.infer_table_definition(header, rows) == [
        ("id", "bigint"),
        ("price", "double precision"),
        ("ordered", "timestamp"),
        ("paid", "boolean"),
        ("note", "text"),
        ("empty", "text"),
    ]


def test_load_sheet_streams_rows_in_batches_and_commits_once():
    connection = FakeConnection()
    result = SheetResult("Sheet1")
    rows = [HEADER, *(order_row(number) for number in range(1, 6))]

    task_one.load_sheet(connection, rows, make_settings(), "part.xlsx", result)

    assert [len(batch) for batch in connection.batches] == [2, 2, 1]
    assert connection.batches[0][0] == [
        1,
        "Customer 1",
        1.5,
        dt.datetime(2024, 1, 2),
    ]
    assert connection.commits == 1
    assert connection.checkpoints == [("orders", "part.xlsx", "Sheet1", 5, 0)]
    assert (result.status, result.rows_loaded, result.batches) == ("loaded", 5, 3)


def test_load_sheet_rejects_invalid_rows_and_loads_the_rest():
    connection = FakeConnection()
    result = SheetResult("Sheet1")
    rows = [
        [None, None],
        HEADER,
        order_row(1),
        ["not a number", "Bob", 1, None],
        [],
        [None, "Nobody", 2, None],
        order_row(2),
    ]

    task_one.load_sheet(connection, rows, make_settings(), "part.xlsx", result)

    assert [row[0] for batch in connection.batches for row in batch] == [1, 2]
    assert (result.rows_loaded, result.rows_rejected) == (2, 2)
    assert connection.checkpoints == [("orders", "part.xlsx", "Sheet1", 2, 2)]


def test_load_sheet_marks_sheet_without_rows_as_empty():
    connection = FakeConnection()
    result = SheetResult("Empty")

    task_one.load_sheet(connection, [(), (None, "")], make_settings(), "x", result)

    assert result.status == "empty"
    assert connection.commits == 0


def test_load_sheet_with_header_only_records_a_checkpoint():
    connection = FakeConnection()
    result = SheetResult("Header")

    task_one.load_sheet(connection, [HEADER], make_settings(), "x.xlsx", result)

    assert (result.status, result.rows_loaded, result.batches) == ("loaded", 0, 0)
    assert connection.checkpoints == [("orders", "x.xlsx", "Header", 0, 0)]


@pytest.mark.parametrize("engine", ENGINES)
def test_process_workbook_loads_every_sheet(tmp_path, database, engine):
    path = write_workbook(
        tmp_path / "orders.xlsx",
        {
            "January": [HEADER, order_row(1), order_row(2), order_row(3)],
            "Empty": [],
            "HeaderOnly": [HEADER],
            "February": [[None], HEADER, order_row(4)],
        },
    )

    outcome = task_one.process_workbook(make_task(path), make_settings(engine=engine))

    assert outcome.error == ""
    assert [(s.sheet, s.status, s.rows_loaded) for s in outcome.sheets] == [
        ("January", "loaded", 3),
        ("Empty", "empty", 0),
        ("HeaderOnly", "loaded", 0),
        ("February", "loaded", 1),
    ]
    assert database.rows == [
        [n, f"Customer {n}", n * 1.5, dt.datetime(2024, 1, 1) + dt.timedelta(days=n)]
        for n in range(1, 5)
    ]
    assert [params[2] for params in database.checkpoints] == [
        "January",
        "HeaderOnly",
        "February",
    ]
    assert len(database.connections) == 1
    assert database.connections[0].closed


def test_process_workbook_skips_sheets_loaded_by_an_earlier_run(tmp_path, database):
    path = write_workbook(
        tmp_path / "orders.xlsx",
        {"Done": [HEADER, order_row(1)], "Todo": [HEADER, order_row(2)]},
    )

    outcome = task_one.process_workbook(make_task(path, {"Done"}), make_settings())

    assert [(s.sheet, s.status) for s in outcome.sheets] == [
        ("Done", "skipped"),
        ("Todo", "loaded"),
    ]
    assert database.rows == [order_row(2)]


def test_process_workbook_continues_after_a_sheet_with_a_bad_header(tmp_path, database):
    path = write_workbook(
        tmp_path / "orders.xlsx",
        {"Bad": [["order_id", "unknown"], [1, 2]], "Good": [HEADER, order_row(1)]},
    )

    outcome = task_one.process_workbook(make_task(path), make_settings())

    bad, good = outcome.sheets
    assert bad.status == "failed"
    assert "not in the target table: unknown" in bad.error
    assert good.status == "loaded"
    assert len(database.rows) == 1


def test_process_workbook_retries_after_a_lost_connection(tmp_path, database):
    path = write_workbook(tmp_path / "orders.xlsx", {"Sheet1": [HEADER, order_row(1)]})
    database.failures.append((1, psycopg.OperationalError("server closed")))

    outcome = task_one.process_workbook(make_task(path), make_settings())

    assert outcome.sheets[0].status == "loaded"
    assert len(database.connections) == 2
    assert database.connections[0].batches == []
    assert database.connections[1].batches == [[order_row(1)]]


def test_process_workbook_gives_up_after_the_configured_retries(tmp_path, database):
    path = write_workbook(
        tmp_path / "orders.xlsx", {"Sheet1": [HEADER, *map(order_row, range(1, 6))]}
    )
    database.failures.extend(
        (3, psycopg.OperationalError("server closed")) for _ in range(3)
    )

    outcome = task_one.process_workbook(make_task(path), make_settings(retries=2))

    result = outcome.sheets[0]
    assert result.status == "failed"
    assert "OperationalError" in result.error
    assert (result.batches, result.failed_batches, result.rows_loaded) == (0, 3, 0)
    assert len(database.connections) == 3
    assert database.rows == []


def test_process_workbook_does_not_retry_data_errors(tmp_path, database):
    path = write_workbook(tmp_path / "orders.xlsx", {"Sheet1": [HEADER, order_row(1)]})
    database.failures.append(
        (1, psycopg.errors.InvalidTextRepresentation("invalid input syntax"))
    )

    outcome = task_one.process_workbook(make_task(path), make_settings())

    assert outcome.sheets[0].status == "failed"
    assert outcome.sheets[0].failed_batches == 1
    assert len(database.connections) == 1


@pytest.mark.parametrize("engine", ENGINES)
def test_process_workbook_reports_an_unreadable_file(tmp_path, database, engine):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"this is not a zip archive")

    outcome = task_one.process_workbook(make_task(path), make_settings(engine=engine))

    assert outcome.error
    assert outcome.sheets == []
    assert database.connections == []


def test_run_totals_aggregate_workbook_results():
    totals = RunTotals(files_total=3)
    totals.add(
        WorkbookResult(
            "a.xlsx",
            100,
            [
                SheetResult("s1", rows_loaded=10, rows_rejected=1, batches=2),
                SheetResult("s2", status="failed", failed_batches=1, error="boom"),
                SheetResult("s3", status="empty"),
            ],
        )
    )
    totals.add(WorkbookResult("b.xlsx", 50, [SheetResult("s1", status="skipped")]))
    totals.add(WorkbookResult("c.xlsx", 25, error="BadZipFile"))

    assert (totals.files_processed, totals.files_skipped, totals.files_failed) == (
        1,
        1,
        1,
    )
    assert (totals.sheets_processed, totals.sheets_failed, totals.sheets_skipped) == (
        2,
        1,
        1,
    )
    assert (totals.rows_loaded, totals.rows_rejected) == (10, 1)
    assert (totals.batches, totals.failed_batches) == (2, 1)
    assert (totals.bytes_processed, totals.bytes_seen) == (100, 175)


@pytest.mark.parametrize(
    "seconds, expected",
    [
        (0, "00h 00m 00.00s"),
        (151.42, "00h 02m 31.42s"),
        (3725.5, "01h 02m 05.50s"),
        (200_000, "2d 07h 33m 20.00s"),
    ],
)
def test_format_duration(seconds, expected):
    assert task_one.format_duration(seconds) == expected


@pytest.mark.parametrize(
    "count, expected",
    [
        (512, "512.00 B"),
        (6_324_049, "6.32 MB"),
        (1.5 * 10**9, "1.50 GB"),
        (2 * 10**12, "2.00 TB"),
        (3 * 10**15, "3.00 PB"),
    ],
)
def test_format_bytes_uses_decimal_units(count, expected):
    assert task_one.format_bytes(count) == expected


def test_rate_is_zero_when_no_time_has_passed():
    assert task_one.rate(1000, 0) == 0.0
    assert task_one.rate(1000, 4) == 250.0


def make_summary(bytes_processed, rows_loaded, seconds, growth=None):
    started = dt.datetime(2026, 1, 1, 9, 0, tzinfo=dt.timezone.utc)
    totals = RunTotals(
        files_total=4,
        files_processed=4,
        sheets_processed=8,
        rows_loaded=rows_loaded,
        batches=10,
        bytes_processed=bytes_processed,
    )
    return BenchmarkSummary(
        started_at=started,
        finished_at=started + dt.timedelta(seconds=seconds),
        seconds=seconds,
        totals=totals,
        table="public.orders",
        engine="openpyxl",
        workers=2,
        batch_size=50_000,
        table_growth=growth,
    )


def test_summary_reports_measured_values_and_a_labelled_projection():
    summary = make_summary(
        bytes_processed=10**9, rows_loaded=5_000_000, seconds=100, growth=3 * 10**9
    )

    text = task_one.format_summary(summary)

    assert summary.bytes_per_second == 10**7
    assert summary.rows_per_second == 50_000
    assert "Start Time        : 2026-01-01 09:00:00+00:00" in text
    assert "End Time          : 2026-01-01 09:01:40+00:00" in text
    assert "Execution Time    : 00h 01m 40.00s (100.00 seconds)" in text
    assert "Rows Processed    : 5,000,000" in text
    assert "Data Processed    : 1.00 GB (1,000,000,000 bytes)" in text
    assert "Throughput        : 10.00 MB/s" in text
    assert "Rows/Second       : 50,000" in text
    assert "Successful Batches: 10" in text
    assert "Failed Batches    : 0" in text
    assert "2 TB figures below are PROJECTED" in text
    assert "Tested Dataset    : 1.00 GB (0.0500% of 2 TB) - Actual" in text
    assert "2 TB Load Time    : 2d 07h 33m 20.00s at 10.00 MB/s - Projected" in text
    assert "2 TB Row Count    : ~10,000,000,000 - Projected" in text
    assert "2 TB Table Size   : ~6.00 TB in PostgreSQL - Projected" in text


def test_summary_of_a_full_2tb_run_is_reported_as_actual():
    text = task_one.format_summary(make_summary(2 * 10**12, 10**10, 86_400))

    assert "RESULT: ACTUAL 2 TB RUN" in text
    assert "Projected" not in text


def test_summary_without_processed_data_has_no_projection():
    summary = make_summary(0, 0, 0)

    text = task_one.format_summary(summary)

    assert "Throughput        : 0.00 B/s" in text
    assert "not available, no data was processed" in text
    assert task_one.summary_record(summary)["projected_2tb_seconds"] is None


def test_summary_record_is_json_serialisable():
    record = task_one.summary_record(make_summary(10**9, 1000, 10))

    restored = json.loads(json.dumps(record))
    assert restored["execution_seconds"] == 10
    assert restored["throughput_bytes_per_second"] == 10**8
    assert restored["projected_2tb_seconds"] == 20_000
    assert restored["full_2tb_processed"] is False


def test_parse_args_reads_connection_settings_from_pg_environment(monkeypatch):
    monkeypatch.setenv("PGHOST", "db.internal")
    monkeypatch.setenv("PGPORT", "6543")
    monkeypatch.setenv("PGDATABASE", "warehouse")
    monkeypatch.setenv("PGUSER", "loader")

    args = task_one.parse_args(["--input-dir", "data", "--table", "sales.orders"])

    assert (args.host, args.port, args.dbname, args.user) == (
        "db.internal",
        6543,
        "warehouse",
        "loader",
    )
    assert args.table == ("sales", "orders")
    assert args.batch_size == 50_000
    assert args.engine == "openpyxl"
    assert args.workers >= 1
    settings = task_one.connection_settings(args)
    assert "password" not in settings


@pytest.mark.parametrize(
    "extra",
    [
        ["--batch-size", "0"],
        ["--workers", "0"],
        ["--retries", "-1"],
        ["--engine", "pandas"],
        ["--table", "a.b.c"],
        ["--table", ".orders"],
    ],
)
def test_parse_args_rejects_invalid_values(extra):
    with pytest.raises(SystemExit):
        task_one.parse_args(["--input-dir", "data", "--table", "orders", *extra])


def test_main_fails_cleanly_for_a_missing_input_directory(tmp_path, run_main, capsys):
    code = run_main(["--input-dir", str(tmp_path / "missing"), "--table", "orders"])

    assert code == 1
    assert "Input directory not found" in capsys.readouterr().out


def test_main_fails_cleanly_when_no_workbooks_are_found(tmp_path, run_main, capsys):
    code = run_main(["--input-dir", str(tmp_path), "--table", "orders"])

    assert code == 1
    assert "No files matching" in capsys.readouterr().out


def test_main_measures_the_run_and_writes_a_report(
    tmp_path, monkeypatch, database, run_main, capsys
):
    input_dir = tmp_path / "data"
    write_workbook(
        input_dir / "part1.xlsx",
        {"Jan": [HEADER, order_row(1), order_row(2)], "Feb": [HEADER, order_row(3)]},
    )
    write_workbook(input_dir / "more" / "part2.xlsx", {"Mar": [HEADER, order_row(4)]})
    monkeypatch.setattr(task_one, "prepare_target", lambda *args: list(ORDER_COLUMNS))
    monkeypatch.setattr(task_one, "completed_sheets", lambda *args: {})
    sizes = iter([1000, 9000])
    monkeypatch.setattr(task_one, "table_size", lambda *args: next(sizes))
    report = tmp_path / "report.json"

    code = run_main(
        [
            "--input-dir",
            str(input_dir),
            "--table",
            "orders",
            "--workers",
            "1",
            "--batch-size",
            "1",
            "--report-file",
            str(report),
        ]
    )

    assert code == 0
    record = json.loads(report.read_text(encoding="utf-8"))
    assert record["files_processed"] == 2
    assert record["sheets_processed"] == 3
    assert record["rows_loaded"] == 4
    assert record["batches"] == 4
    assert record["failed_batches"] == 0
    assert record["table_growth_bytes"] == 8000
    total_bytes = sum(path.stat().st_size for path in input_dir.rglob("*.xlsx"))
    assert record["bytes_processed"] == total_bytes
    assert record["execution_seconds"] > 0
    start = dt.datetime.fromisoformat(record["start_time"])
    end = dt.datetime.fromisoformat(record["end_time"])
    assert start <= end
    assert record["rows_per_second"] == pytest.approx(
        4 / record["execution_seconds"], rel=0.01
    )
    assert sorted(row[0] for row in database.rows) == [1, 2, 3, 4]
    output = capsys.readouterr().out
    assert "EXCEL TO POSTGRESQL BENCHMARK" in output
    assert "Rows Processed    : 4" in output
    assert "Files Processed   : 2 of 2" in output


def test_main_returns_an_error_code_when_a_sheet_fails(
    tmp_path, monkeypatch, database, run_main
):
    write_workbook(tmp_path / "bad.xlsx", {"Sheet1": [["wrong", "header"], [1, 2]]})
    monkeypatch.setattr(task_one, "prepare_target", lambda *args: list(ORDER_COLUMNS))
    monkeypatch.setattr(task_one, "completed_sheets", lambda *args: {})
    monkeypatch.setattr(task_one, "table_size", lambda *args: 0)

    code = run_main(
        ["--input-dir", str(tmp_path), "--table", "orders", "--workers", "1"]
    )

    assert code == 1


@pytest.mark.parametrize(
    "text, expected",
    [
        ("500MB", 500 * 10**6),
        ("1.5 GB", 1_500_000_000),
        ("2TB", 2 * 10**12),
        ("2tb", 2 * 10**12),
    ],
)
def test_generator_parses_sizes(text, expected):
    assert generate_test_data.parse_size(text) == expected


@pytest.mark.parametrize("text", ["2 TiB", "lots", "-1GB", ""])
def test_generator_rejects_invalid_sizes(text):
    with pytest.raises(argparse.ArgumentTypeError):
        generate_test_data.parse_size(text)


def test_generator_rejects_more_rows_than_an_excel_sheet_can_hold(tmp_path):
    with pytest.raises(SystemExit):
        generate_test_data.parse_args(
            ["--output-dir", str(tmp_path), "--rows-per-sheet", "1048576"]
        )


def test_generated_workbook_has_the_requested_sheets_and_rows(tmp_path):
    path = tmp_path / "dataset.xlsx"

    generate_test_data.write_workbook(path, sheets=3, rows=20, first_id=101, seed=1)

    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames == ["Sheet1", "Sheet2", "Sheet3"]
        ids = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            assert list(rows[0]) == generate_test_data.HEADER
            assert len(rows) == 21
            ids += [row[0] for row in rows[1:]]
    finally:
        workbook.close()
    assert ids == list(range(101, 161))


@pytest.mark.parametrize("engine", ENGINES)
def test_loader_reads_generated_workbooks(tmp_path, database, engine):
    path = tmp_path / "dataset.xlsx"
    generate_test_data.write_workbook(path, sheets=2, rows=30, first_id=1, seed=7)
    header, sample = task_one.read_sample([path], engine)
    columns = (
        Column("order_id", "int8", not_null=True),
        Column("customer_name", "text"),
        Column("region", "text"),
        Column("product", "text"),
        Column("quantity", "int4"),
        Column("unit_price", "numeric"),
        Column("order_date", "timestamp"),
    )

    outcome = task_one.process_workbook(
        make_task(path), make_settings(columns=columns, engine=engine, batch_size=7)
    )

    assert task_one.infer_table_definition(header, sample) == [
        ("order_id", "bigint"),
        ("customer_name", "text"),
        ("region", "text"),
        ("product", "text"),
        ("quantity", "bigint"),
        ("unit_price", "double precision"),
        ("order_date", "timestamp"),
    ]
    assert [s.rows_loaded for s in outcome.sheets] == [30, 30]
    assert [s.rows_rejected for s in outcome.sheets] == [0, 0]
    assert [s.batches for s in outcome.sheets] == [5, 5]
    assert [row[0] for row in database.rows] == list(range(1, 61))
    assert all(isinstance(row[6], dt.datetime) for row in database.rows)


def test_generator_copies_workbooks_until_the_target_size(tmp_path, capsys):
    code = generate_test_data.main(
        [
            "--output-dir",
            str(tmp_path),
            "--files",
            "1",
            "--sheets-per-file",
            "2",
            "--rows-per-sheet",
            "10",
            "--workers",
            "1",
            "--target-size",
            "20KB",
        ]
    )

    assert code == 0
    files = sorted(tmp_path.glob("dataset_*.xlsx"))
    sizes = [path.stat().st_size for path in files]
    assert len(files) > 1
    assert sum(sizes) >= 20_000 > sum(sizes[:-1])
    assert len({path.read_bytes() for path in files}) == 1
    output = capsys.readouterr().out
    assert f"Dataset: {len(files)} files, {len(files) * 2} sheets" in output
    assert "repeat the order_id values" in output


def test_generator_refuses_to_mix_with_an_existing_dataset(tmp_path):
    (tmp_path / "dataset_00001.xlsx").write_bytes(b"existing")

    code = generate_test_data.main(["--output-dir", str(tmp_path), "--workers", "1"])

    assert code == 1


def test_generator_checks_free_disk_space_before_copying(tmp_path, monkeypatch):
    class Usage:
        free = 10**6

    monkeypatch.setattr(generate_test_data.shutil, "disk_usage", lambda path: Usage)

    code = generate_test_data.main(
        [
            "--output-dir",
            str(tmp_path),
            "--files",
            "1",
            "--sheets-per-file",
            "1",
            "--rows-per-sheet",
            "5",
            "--workers",
            "1",
            "--target-size",
            "2TB",
        ]
    )

    assert code == 1
    assert len(list(tmp_path.glob("*.xlsx"))) == 1


@pytest.fixture
def postgres_schema():
    schema = f"task_one_test_{uuid.uuid4().hex[:8]}"
    with psycopg.connect(autocommit=True) as connection:
        connection.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))
    yield schema
    with psycopg.connect(autocommit=True) as connection:
        connection.execute(
            sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(schema))
        )


@requires_postgres
@pytest.mark.parametrize("engine", ENGINES)
def test_end_to_end_load_into_postgres(tmp_path, postgres_schema, engine, run_main):
    input_dir = tmp_path / "data"
    write_workbook(
        input_dir / "part1.xlsx",
        {"Jan": [HEADER, *map(order_row, range(1, 6))], "Empty": []},
    )
    write_workbook(
        input_dir / "part2.xlsx",
        {"Feb": [HEADER, *map(order_row, range(6, 9)), ["bad id", "Eve", 1, None]]},
    )
    table = sql.Identifier(postgres_schema, "orders")
    report = tmp_path / "report.json"
    args = [
        "--input-dir",
        str(input_dir),
        "--table",
        f"{postgres_schema}.orders",
        "--workers",
        "2",
        "--batch-size",
        "2",
        "--engine",
        engine,
        "--report-file",
        str(report),
    ]

    def loaded():
        with psycopg.connect() as connection:
            return connection.execute(
                sql.SQL("SELECT count(*), sum(order_id) FROM {}").format(table)
            ).fetchone()

    assert run_main(args) == 0
    record = json.loads(report.read_text(encoding="utf-8"))
    assert (record["rows_loaded"], record["rows_rejected"]) == (8, 1)
    assert record["table_growth_bytes"] > 0
    assert loaded() == (8, 36)
    with psycopg.connect() as connection:
        types = dict(
            connection.execute(
                "SELECT column_name, data_type FROM information_schema.columns "
                "WHERE table_schema = %s AND table_name = 'orders'",
                (postgres_schema,),
            ).fetchall()
        )
    assert types == {
        "order_id": "bigint",
        "customer": "text",
        "amount": "double precision",
        "ordered_at": "timestamp without time zone",
    }

    assert run_main(args) == 0
    record = json.loads(report.read_text(encoding="utf-8"))
    assert (record["rows_loaded"], record["sheets_skipped"]) == (0, 2)
    assert loaded() == (8, 36)

    assert run_main([*args, "--truncate"]) == 0
    assert json.loads(report.read_text(encoding="utf-8"))["rows_loaded"] == 8
    assert loaded() == (8, 36)
